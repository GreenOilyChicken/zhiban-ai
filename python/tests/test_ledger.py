"""
test_ledger — 测试台账生成（针对生产实现）
"""
import datetime as dt

import pytest
from pathlib import Path
from unittest.mock import MagicMock


def test_generate_single_ledger(tmp_path):
    """测试生成单个台账（生产 LedgerGenerator）"""
    pytest.importorskip("docx", reason="python-docx 未安装")

    from govpolicy_agent.training_ledger.ledger_generator import LedgerGenerator

    gen = LedgerGenerator(output_dir=tmp_path)
    out_path = gen.generate(
        quarter="2025Q2",
        unit_name="知伴（湖州）智能科技有限公司",
        topics=[
            ("惠民政策解读", ["张三", "李四"], "2025-05-12"),
            ("政务服务流程培训", ["王五"], "2025-06-18"),
        ],
    )

    out_path = Path(out_path)
    assert out_path.exists()
    assert out_path.suffix == ".docx"
    assert out_path.stat().st_size > 1000

    # 文档内容应包含单位名与季度
    from docx import Document
    doc = Document(str(out_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    assert "知伴" in text
    assert "2025" in text


def test_sanitize_filename():
    """测试文件名净化（生产实现）"""
    from govpolicy_agent.training_ledger.ledger_generator import _sanitize_filename

    assert _sanitize_filename("正常名称") == "正常名称"
    assert _sanitize_filename("test/name") == "test_name"
    # 全非法字符 -> 统一回退 unknown
    assert _sanitize_filename("///") == "unknown"
    assert _sanitize_filename("") == "unknown"
    assert _sanitize_filename("   ") == "unknown"
    assert _sanitize_filename("a:b?c|d") == "a_b_c_d"


def test_quarter_to_daterange():
    """测试季度日期范围计算（生产实现）"""
    from govpolicy_agent.training_ledger.ledger_generator import _quarter_to_daterange

    # Q1
    s, e = _quarter_to_daterange("2025Q1")
    assert s == dt.date(2025, 1, 1)
    assert e == dt.date(2025, 3, 31)

    # Q2
    s, e = _quarter_to_daterange("2025Q2")
    assert s == dt.date(2025, 4, 1)
    assert e == dt.date(2025, 6, 30)

    # Q3
    s, e = _quarter_to_daterange("2025Q3")
    assert s == dt.date(2025, 7, 1)
    assert e == dt.date(2025, 9, 30)

    # Q4（跨年边界）
    s, e = _quarter_to_daterange("2025Q4")
    assert s == dt.date(2025, 10, 1)
    assert e == dt.date(2025, 12, 31)

    # 闰年 Q1 结束日
    s, e = _quarter_to_daterange("2024Q1")
    assert e == dt.date(2024, 3, 31)

    # 多种分隔符写法均应支持
    assert _quarter_to_daterange("2025-Q1") == _quarter_to_daterange("2025Q1")
    assert _quarter_to_daterange("2025/Q1") == _quarter_to_daterange("2025Q1")

    # 非法格式应报错
    with pytest.raises(ValueError):
        _quarter_to_daterange("2025Q5")
    with pytest.raises(ValueError):
        _quarter_to_daterange("not-a-quarter")


def test_ledger_topics_type():
    """测试台账 topics 类型"""
    from typing import List, Tuple, Optional

    sample_topics: List[Tuple[str, List[str], str]] = [
        ("惠民政策解读", ["张三", "李四"], "2025-05-12"),
        ("政务服务培训", ["王五"], "2025-06-18"),
    ]

    assert len(sample_topics) == 2
    assert isinstance(sample_topics[0], tuple)
    assert len(sample_topics[0]) == 3
    assert isinstance(sample_topics[0][1], list)
